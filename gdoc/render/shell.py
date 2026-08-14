"""Turn the master template into a filled "shell": front matter completed,
sample body removed, ready for generated content to be spliced in.

Everything the eye recognises as the Altery house document lives in the front
matter and the page furniture, so none of it is rebuilt here. The template is
copied and edited in place, which is what keeps the output pixel-identical to
the master.
"""

import copy
import re
import shutil
import zipfile

from docx import Document

from gdoc.render.ooxml import (BODY_SZ, HEADING_EMPHASIS, HEADING_LINE_SPACING,
                               HEADING_SPACE_AFTER, HEADING_SPACE_BEFORE, USABLE_TWIPS,
                               el, qn, set_fonts, set_ppr_child, sub)

# The template's own "highlight as appropriate" mechanism for the Document
# Classification table is cell shading, not a run highlight. Each class has
# its own colour on the label cell; marking a class means painting the
# description cell to match.
CLASS_FILLS = {
    "Confidential (C)": "f4cccc",
    "Restricted (R)": "fce5cd",
    "Internal (I)": "fff2cc",
    "Public (P)": "d9ead3",
}

VERSION_CONTROL_LABELS = {
    "Document Owner": "owner",
    "Date of Last Approval": "last_approval",
    "Review Frequency": "review_frequency",
    "Board Ratification Date": "board_ratification",
    "Policy Distribution": "distribution",
}

TITLE_PLACEHOLDER = "(Name of) Framework/Policy"
VERSION_PLACEHOLDER = "Version: 1.0"
DATE_PLACEHOLDER = "May 2025"
RUNNING_HEAD_PLACEHOLDER = "Altery - xxx Policy"

BLACK = "000000"


class TemplateError(RuntimeError):
    """The master template is not shaped the way this code expects."""


# --------------------------------------------------------------- text edits --
def _clear_placeholder_marks(run):
    """Strip the yellow highlight and the red instruction colour from a run.

    In the master, yellow marks "a human must fill this in" and red marks
    "this is guidance, not content". Once we have written a real value both
    marks are actively misleading, so they go.
    """
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return
    for highlight in rpr.findall(qn("w:highlight")):
        rpr.remove(highlight)
    for color in rpr.findall(qn("w:color")):
        color.set(qn("w:val"), BLACK)


def replace_in_paragraph(paragraph, old, new):
    """Replace `old` with `new` even when it is split across runs.

    Google Docs starts a new run wherever formatting changes, so
    "(Name of) Framework/Policy" is really two runs (yellow "(Name of) " plus
    red "Framework/Policy") and "Version: 1.0" is "Version: " plus a
    highlighted "1.0". Matching run by run therefore finds nothing. Match the
    joined paragraph text instead, write the replacement into the first run
    that overlaps the match, and blank the overlap out of the rest. Writing
    into an existing run is what preserves the font, size and centring.
    """
    runs = paragraph.runs
    if not runs:
        return False
    joined = "".join(r.text for r in runs)
    start = joined.find(old)
    if start < 0:
        return False
    end = start + len(old)

    position = 0
    written = False
    for run in runs:
        run_start, run_end = position, position + len(run.text)
        position = run_end
        if run_end <= start or run_start >= end:
            continue
        head = run.text[: max(0, start - run_start)]
        tail = run.text[max(0, end - run_start):] if run_end > end else ""
        if not written:
            run.text = head + new + tail
            _clear_placeholder_marks(run)
            written = True
        else:
            run.text = head + tail
    return True


def _paragraph_text(paragraph):
    return "".join(r.text for r in paragraph.runs)


def _delete(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


# ------------------------------------------------------------------- cover ---
def fill_cover(doc, meta):
    """Fill the cover block and drop the unused alternate-title lines.

    The master offers two title lines separated by the word "or", so an author
    can show a Framework and the Policy it sits under. Most documents have one
    title, and leaving a stray "or" on the cover looks like a defect, so both
    extra lines are removed unless `alt_title` is set.
    """
    title_paragraphs = [p for p in doc.paragraphs
                        if TITLE_PLACEHOLDER in _paragraph_text(p)]
    if len(title_paragraphs) < 2:
        raise TemplateError(
            f"expected two {TITLE_PLACEHOLDER!r} cover lines, "
            f"found {len(title_paragraphs)}"
        )

    replace_in_paragraph(title_paragraphs[0], TITLE_PLACEHOLDER, meta["cover_title"])

    if meta["cover_alt_title"]:
        replace_in_paragraph(title_paragraphs[1], TITLE_PLACEHOLDER,
                             meta["cover_alt_title"])
    else:
        _delete(title_paragraphs[1])
        for paragraph in doc.paragraphs:
            if _paragraph_text(paragraph).strip() == "or":
                _delete(paragraph)
                break

    for paragraph in doc.paragraphs:
        text = _paragraph_text(paragraph)
        if VERSION_PLACEHOLDER in text:
            replace_in_paragraph(paragraph, VERSION_PLACEHOLDER,
                                 f"Version: {meta['version']}")
        elif DATE_PLACEHOLDER in text:
            replace_in_paragraph(paragraph, DATE_PLACEHOLDER, meta["date"])


# ------------------------------------------------------------------ tables ---
def _apply_mark_format(run_element, paragraph_element):
    """Give a run the paragraph mark's font and size where it has none of its own.

    An empty template cell still holds a bare `<w:r/>` carrying no `w:rPr`. Text
    written into it would render at the document default rather than the
    template's, which is visibly smaller. The intended formatting is on the
    paragraph mark (`w:pPr/w:rPr`), so that is what gets copied down.
    """
    r_pr = run_element.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = el("rPr")
        run_element.insert(0, r_pr)

    p_pr = paragraph_element.find(qn("w:pPr"))
    mark = p_pr.find(qn("w:rPr")) if p_pr is not None else None

    for tag, fallback in (("w:sz", BODY_SZ), ("w:szCs", BODY_SZ)):
        if r_pr.find(qn(tag)) is not None:
            continue
        source = mark.find(qn(tag)) if mark is not None else None
        value = source.get(qn("w:val")) if source is not None else fallback
        sub(r_pr, tag.split(":")[1], val=value)

    if r_pr.find(qn("w:rFonts")) is None:
        source = mark.find(qn("w:rFonts")) if mark is not None else None
        if source is not None:
            r_pr.insert(0, copy.deepcopy(source))
        else:
            set_fonts(r_pr)


def _set_cell_text(cell, text):
    """Write `text` into a cell, keeping the first run's formatting.

    Reusing the existing run rather than making a new one inherits the
    template's font, size and alignment for free. Surplus runs and paragraphs
    are dropped so nothing of the sample data survives.
    """
    paragraphs = cell.paragraphs
    target = paragraphs[0]
    for extra in paragraphs[1:]:
        _delete(extra)

    runs = target.runs
    if not runs:
        sub(target._element, "r")
        runs = target.runs

    runs[0].text = text
    _clear_placeholder_marks(runs[0])
    _apply_mark_format(runs[0]._element, target._element)
    for extra in runs[1:]:
        extra._element.getparent().remove(extra._element)


def fill_version_control(table, meta):
    """Fill the five-row Document Owner / approval-dates table."""
    for row in table.rows:
        label = row.cells[0].text.strip()
        key = VERSION_CONTROL_LABELS.get(label)
        if key and meta.get(key):
            _set_cell_text(row.cells[1], meta[key])


def fill_revisions(table, revisions):
    """Rebuild the revision-history rows from front matter.

    Row 0 is the header and row 1 is the sample row, which doubles as the
    formatting prototype for however many revisions the author declared.
    """
    if not revisions:
        return
    prototype = copy.deepcopy(table.rows[1]._tr)

    for row in list(table.rows[1:]):
        row._tr.getparent().remove(row._tr)

    for revision in revisions:
        new_row = copy.deepcopy(prototype)
        table._tbl.append(new_row)
        cells = table.rows[-1].cells
        values = (revision["version"], revision["date"], revision["author"],
                  revision["approved_by"], revision["approval_date"],
                  revision["section"], revision["change"])
        for cell, value in zip(cells, values):
            _set_cell_text(cell, value)


def mark_classification(table, classification):
    """Shade the chosen classification row the way the master does it."""
    for row in table.rows[1:]:
        label = row.cells[0].text.strip()
        description = row.cells[1]
        tc_pr = description._tc.find(qn("w:tcPr"))
        if tc_pr is None:
            continue
        for shd in tc_pr.findall(qn("w:shd")):
            tc_pr.remove(shd)
        if label == classification:
            shd = el("shd", val="clear", color="auto",
                     fill=CLASS_FILLS.get(classification, "fff2cc"))
            tc_pr.append(shd)


# -------------------------------------------------------------------- body ---
def add_page_break_before(doc, text):
    """Start the paragraph whose text is `text` on a new page.

    w:pageBreakBefore has to sit where the schema wants it or Word rejects the
    whole w:pPr, which is what `set_ppr_child` is for.
    """
    for paragraph in doc.paragraphs:
        if _paragraph_text(paragraph).strip() != text:
            continue
        p_pr = paragraph._element.find(qn("w:pPr"))
        if p_pr is None:
            p_pr = el("pPr")
            paragraph._element.insert(0, p_pr)
        set_ppr_child(p_pr, "w:pageBreakBefore", val="1")
        return True
    return False


def remove_empty_paragraphs_before(doc, text):
    """Drop blank filler paragraphs sitting immediately above `text`.

    The master pads its sections apart with empty paragraphs. Once the section
    starts on a page of its own those spacers have nothing to space, and if
    they happen to cross the page boundary they produce an entirely blank page
    ahead of the break.
    """
    paragraphs = doc.paragraphs
    target = next((i for i, p in enumerate(paragraphs)
                   if _paragraph_text(p).strip() == text), None)
    if target is None:
        return 0
    removed = 0
    for paragraph in reversed(paragraphs[:target]):
        if _paragraph_text(paragraph).strip():
            break
        if paragraph._element.find(qn("w:tbl")) is not None:
            break
        _delete(paragraph)
        removed += 1
    return removed


HEADING_STYLE_RE = re.compile(r"Heading([1-6])$")


def _heading_styles(doc):
    """Yield (level, style) for each Heading1..Heading6 style definition."""
    for style in doc.styles.element.findall(qn("w:style")):
        match = HEADING_STYLE_RE.fullmatch(style.get(qn("w:styleId")) or "")
        if match:
            yield int(match.group(1)), style


def _style_ppr(style):
    """The style's w:pPr, created in the right place if it has none.

    w:pPr must precede w:rPr inside w:style or Word rejects the style.
    """
    p_pr = style.find(qn("w:pPr"))
    if p_pr is not None:
        return p_pr
    p_pr = el("pPr")
    r_pr = style.find(qn("w:rPr"))
    if r_pr is None:
        style.append(p_pr)
    else:
        r_pr.addprevious(p_pr)
    return p_pr


def _style_rpr(style):
    """The style's w:rPr, created at the end if it has none."""
    r_pr = style.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = el("rPr")
        style.append(r_pr)
    return r_pr


def _set_rpr_flag(r_pr, tag, on):
    """Set or clear a boolean run property such as w:b or w:i.

    Removing the element is not enough on its own for a style that inherits the
    flag, so it is written as val="0" rather than dropped.
    """
    for pair in (tag, f"{tag}Cs"):
        for existing in r_pr.findall(qn(f"w:{pair}")):
            r_pr.remove(existing)
        sub(r_pr, pair, val="1" if on else "0")


def normalise_heading_emphasis(doc):
    """Make weight fall as heading depth grows, for Heading3..Heading6.

    Two defects in the master, and the second is the one readers notice:

    - Heading4..Heading6 inherit Normal's 11pt while body text is 12pt, so a
      fourth-level heading renders *smaller* than the prose underneath it.
    - Heading3 is regular while Heading4 and Heading6 are bold. Since all four
      end up at the same 12pt, a child heading looks stronger than its parent.

    Size is set to 12pt across the four, and emphasis is stepped down instead:
    bold, regular, italic, italic. See HEADING_EMPHASIS for why levels 5 and 6
    share a treatment.
    """
    changed = 0
    for level, style in _heading_styles(doc):
        emphasis = HEADING_EMPHASIS.get(level)
        if emphasis is None:
            continue  # Heading1 and Heading2 keep the template's 16pt and 14pt
        r_pr = _style_rpr(style)
        for tag in ("w:sz", "w:szCs"):
            for existing in r_pr.findall(qn(tag)):
                r_pr.remove(existing)
            sub(r_pr, tag.split(":")[1], val=BODY_SZ)
        _set_rpr_flag(r_pr, "b", emphasis["bold"])
        _set_rpr_flag(r_pr, "i", emphasis["italic"])
        changed += 1
    return changed


def normalise_heading_spacing(doc):
    """Give Heading1..Heading6 1.15 line spacing and room beneath them.

    Two defects in the master, both about the gap around a heading:

    - The heading styles carry `lineRule="auto"` with no `w:line`, so they fall
      back to single spacing while the prose around them is 1.15. A heading that
      wraps to a second line is then tighter than the body text under it.
    - Heading1 and Heading2 set `after="0"` and Heading3..Heading6 set nothing,
      which also resolves to zero. The heading therefore sits hard against the
      first line of its own section, so the eye cannot find where a section
      starts.

    Fixed on the style rather than per paragraph so Word, Google Docs and
    LibreOffice agree, and so the contents list inherits it too.
    """
    changed = 0
    for _level, style in _heading_styles(doc):
        set_ppr_child(_style_ppr(style), "w:spacing",
                      before=HEADING_SPACE_BEFORE, after=HEADING_SPACE_AFTER,
                      line=HEADING_LINE_SPACING, lineRule="auto")
        changed += 1
    return changed


def normalise_heading_keeps(doc):
    """Stop a heading being stranded at the foot of a page, or split in two.

    The master sets keepNext on Heading1..Heading4 but not on Heading5 or
    Heading6, and sets keepLines on none of them. So a deep heading can sit
    alone at the bottom of a page with its section overleaf, and a heading long
    enough to wrap can break across the page boundary.
    """
    changed = 0
    for _level, style in _heading_styles(doc):
        p_pr = _style_ppr(style)
        set_ppr_child(p_pr, "w:keepNext", val="1")
        set_ppr_child(p_pr, "w:keepLines", val="1")
        changed += 1
    return changed


def normalise_heading_indents(doc):
    """Flatten the left indent out of the Heading1..Heading6 style definitions.

    The master indents its heading styles (Heading1 left=720, Heading3
    left=2160, both hanging=360), which is a leftover from the Google Docs
    numbered-heading list. A direct `w:ind` on the paragraph clears this for
    Heading1 and Heading2 but is silently ignored for Heading3, which then
    renders 90pt from the margin while its siblings sit flush.

    Fixing the style rather than each paragraph is both the smaller change and
    the portable one: it does not depend on direct formatting winning, so Word
    and Google Docs agree with LibreOffice.
    """
    changed = 0
    for _level, style in _heading_styles(doc):
        set_ppr_child(_style_ppr(style), "w:ind",
                      left="0", right="0", hanging="0", firstLine="0")
        changed += 1
    return changed


def strip_body(doc):
    """Delete everything from the first Heading1 to the end of the body.

    sectPr is deliberately left in place: it carries the page size, margins,
    titlePg flag and the header/footer wiring, so removing it would take the
    logo and the running head with it.
    """
    body = doc.element.body
    children = list(body)
    start = None
    for index, element in enumerate(children):
        if element.tag != qn("w:p"):
            continue
        p_pr = element.find(qn("w:pPr"))
        if p_pr is None:
            continue
        style = p_pr.find(qn("w:pStyle"))
        if style is not None and style.get(qn("w:val")) == "Heading1":
            start = index
            break
    if start is None:
        raise TemplateError("no Heading1 found in the template, so the body "
                            "boundary cannot be located")

    removed = 0
    for element in children[start:]:
        if element.tag == qn("w:sectPr"):
            continue
        body.remove(element)
        removed += 1
    return removed


# ----------------------------------------------------------------- headers ---
def fill_running_head(path, running_head):
    """Patch the running head, which lives in the header parts, not the body."""
    doc = Document(path)
    for section in doc.sections:
        headers = (section.header, section.first_page_header, section.even_page_header)
        for header in headers:
            if header is None:
                continue
            for paragraph in header.paragraphs:
                for run in paragraph.runs:
                    if RUNNING_HEAD_PLACEHOLDER in run.text:
                        run.text = run.text.replace(RUNNING_HEAD_PLACEHOLDER,
                                                    running_head)
    doc.save(path)


# ---------------------------------------------------------------- comments ---
COMMENT_PARTS = ("word/comments.xml", "word/commentsExtended.xml")
COMMENT_TAG_RE = re.compile(
    r"<w:(commentRangeStart|commentRangeEnd)\b[^>]*/>"
    r"|<w:r\b[^>]*>(?:(?!</w:r>).)*?<w:commentReference\b[^>]*/>.*?</w:r>",
    re.DOTALL,
)


def strip_comments(path):
    """Remove the template's tracked comments from the generated document.

    The master carries 46KB of review comments. Left in place they travel into
    every generated document and, once uploaded, reappear as live Google Docs
    comments on a document that is supposed to be freshly issued.

    Done at the package level because it spans four things that must stay
    consistent: the comment parts, their relationships, their content-type
    overrides, and the anchors inside document.xml.
    """
    with zipfile.ZipFile(path) as archive:
        items = [(item, archive.read(item.filename)) for item in archive.infolist()]

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for item, data in items:
            name = item.filename
            if name in COMMENT_PARTS:
                continue
            if name == "word/document.xml":
                data = COMMENT_TAG_RE.sub("", data.decode("utf-8")).encode("utf-8")
            elif name == "[Content_Types].xml":
                text = data.decode("utf-8")
                for part in COMMENT_PARTS:
                    text = re.sub(
                        rf'<Override PartName="/{re.escape(part)}"[^>]*/>', "", text)
                data = text.encode("utf-8")
            elif name == "word/_rels/document.xml.rels":
                text = data.decode("utf-8")
                text = re.sub(
                    r'<Relationship[^>]*Target="comments(Extended)?\.xml"[^>]*/>',
                    "", text)
                data = text.encode("utf-8")
            archive.writestr(item, data)


# --------------------------------------------------------------- entrypoint --
def build_shell(template_path, output_path, meta):
    """Copy the master, fill its front matter, and clear the sample body."""
    shutil.copyfile(template_path, output_path)

    doc = Document(output_path)
    if len(doc.tables) < 3:
        raise TemplateError(
            f"expected at least 3 front-matter tables in the template, "
            f"found {len(doc.tables)}"
        )
    fill_cover(doc, meta)
    fill_version_control(doc.tables[0], meta)
    fill_revisions(doc.tables[1], meta["revisions"])
    mark_classification(doc.tables[2], meta["classification"])
    normalise_heading_indents(doc)
    normalise_heading_emphasis(doc)
    normalise_heading_spacing(doc)
    normalise_heading_keeps(doc)

    # Give the cover, the document-control block and the contents list a page
    # each. The master runs them together, which leaves the Version Control
    # table split across the cover and page 2.
    for label in ("Version Control", "Contents"):
        remove_empty_paragraphs_before(doc, label)
        if not add_page_break_before(doc, label):
            raise TemplateError(
                f"could not find the {label!r} paragraph to break the page before"
            )

    removed = strip_body(doc)
    doc.save(output_path)

    fill_running_head(output_path, meta["running_head"])
    strip_comments(output_path)
    return removed
