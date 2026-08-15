"""Render the Markdown body into the shell using the template's own idioms.

Pandoc is used only as a Markdown parser (`pandoc -t json`), never as a docx
writer. That distinction matters: pandoc's docx writer emits its own styles
(`Compact`, `BodyText`, `FirstParagraph`) which this Google-Docs-exported
template does not define, and LibreOffice responds to a dangling `w:pStyle` by
discarding the entire `w:pPr` around it. Every bullet and every number then
silently disappears. Parsing only, and writing the OOXML ourselves, avoids the
whole class of problem.
"""

import json
import re
import subprocess
from pathlib import Path

from docx.shared import Emu

from gdoc.render.ooxml import (BODY_FONT, BODY_SZ, CELL_BORDER, CELL_MARGIN_H,
                               CELL_MARGIN_V, HANGING, HDR_FILL, HEAD_COLOR, IND_LEFT,
                               LINE_SPACING, MAX_LIST_LEVEL, NO_FILL, ROW_FILL,
                               TBL_BORDER, USABLE_TWIPS, el, qn, set_fonts, sub, text_run)
from gdoc.export import PandocNotFound, find_pandoc

# 1 inch is 1440 twips and 914400 EMU, so one twip is exactly 635 EMU. An image
# is never widened past the text column, and never upscaled past its own size.
TWIP_EMU = 635
IMAGE_MAX_EMU = USABLE_TWIPS * TWIP_EMU
CAPTION_SZ = "20"       # 10pt half-points
CAPTION_COLOR = "595959"

# Headings that name themselves are never auto-numbered; "Appendix 2" would
# otherwise come out as "13-Appendix 2".
UNNUMBERED_HEADING_RE = re.compile(
    r"^\s*(appendix|appendices|addendum|addenda|annex(e|es|ure)?|"
    r"contents|glossary|schedule)\b",
    re.IGNORECASE,
)

# Pandoc extensions: pipe tables for the table syntax, mark for ==highlight==,
# which is the only way to say "a human still has to fill this in".
PANDOC_FORMAT = "markdown+pipe_tables+mark+strikeout+task_lists"


class BodyError(RuntimeError):
    """The Markdown body could not be parsed or rendered."""


# ---------------------------------------------------------------- numbering --
class HeadingNumberer:
    """Produce the template's "1-", "1.1-", "1.1.1-" heading prefixes.

    The master writes these as literal text rather than as Word list numbering,
    so we do the same. Reproducing the template beats being clever here: a real
    numbered-heading list would renumber differently in Word and Google Docs.
    """

    def __init__(self, enabled=True, top_level=1):
        self.enabled = enabled
        # A document whose top heading is "##" should still number from 1, not
        # from "0.1". Treating the shallowest heading present as level 1 makes
        # the output the same either way, which matters when the body is lifted
        # out of a note that already had its own title line.
        self.top_level = top_level
        self.counters = [0] * 9

    def style_level(self, level):
        """Map a Markdown heading level onto a template Heading style.

        The shallowest heading in the document becomes Heading1, so a file
        whose top level is "##" still gets a 16pt Heading1 rather than a 14pt
        Heading2. This is independent of numbering: it applies even when
        numbering is off.
        """
        return max(1, level - self.top_level + 1)

    def prefix(self, level, heading_text):
        if not self.enabled or UNNUMBERED_HEADING_RE.match(heading_text):
            return ""
        index = max(0, level - self.top_level)
        self.counters[index] += 1
        for deeper in range(index + 1, len(self.counters)):
            self.counters[deeper] = 0
        return ".".join(str(c) for c in self.counters[: index + 1]) + "-"


def shallowest_heading_level(blocks, current=None):
    """Find the smallest header level anywhere in the document."""
    for block in blocks:
        tag, content = block["t"], block.get("c")
        if tag == "Header":
            level = content[0]
            current = level if current is None else min(current, level)
        elif tag in ("BulletList", "OrderedList", "BlockQuote", "Div"):
            nested = content[1] if tag in ("OrderedList", "Div") else content
            for item in (nested if tag != "BlockQuote" else [nested]):
                if isinstance(item, list):
                    current = shallowest_heading_level(item, current)
    return current or 1


# ------------------------------------------------------------------ inlines --
def inline_runs(nodes, out=None, bold=False, italic=False, mono=False,
                highlight=None):
    """Flatten a pandoc inline list into run tuples."""
    if out is None:
        out = []
    for node in nodes:
        if not isinstance(node, dict):
            continue
        tag, content = node["t"], node.get("c")
        if tag == "Str":
            out.append((content, bold, italic, mono, highlight))
        elif tag in ("Space", "SoftBreak"):
            out.append((" ", bold, italic, mono, highlight))
        elif tag == "LineBreak":
            out.append(("\n", bold, italic, mono, highlight))
        elif tag == "Strong":
            inline_runs(content, out, True, italic, mono, highlight)
        elif tag in ("Emph", "Underline"):
            inline_runs(content, out, bold, True, mono, highlight)
        elif tag == "Highlighted":
            inline_runs(content, out, bold, italic, mono, "yellow")
        elif tag == "Code":
            out.append((content[1], bold, italic, True, highlight))
        elif tag == "Quoted":
            marks = "“”" if content[0]["t"] == "DoubleQuote" else "‘’"
            out.append((marks[0], bold, italic, mono, highlight))
            inline_runs(content[1], out, bold, italic, mono, highlight)
            out.append((marks[1], bold, italic, mono, highlight))
        elif tag == "Link":
            inline_runs(content[1], out, bold, italic, mono, highlight)
        elif tag in ("Span", "Strikeout", "SmallCaps", "Superscript", "Subscript",
                     "Cite", "Note"):
            inline_runs(content[-1], out, bold, italic, mono, highlight)
        elif tag == "RawInline":
            continue
        elif isinstance(content, list):
            inline_runs(content, out, bold, italic, mono, highlight)
    return out


def add_runs(parent, runs, sz=BODY_SZ, color=None):
    for text, bold, italic, mono, highlight in runs:
        if text == "":
            continue
        text_run(parent, text, bold=bold, italic=italic, mono=mono,
                 color=color, highlight=highlight, sz=sz)


# ----------------------------------------------------------------- builders --
def make_paragraph(runs, sz=BODY_SZ, justify=True, before="0", align=None,
                   color=None, italic=False):
    paragraph = el("p")
    p_pr = sub(paragraph, "pPr")
    sub(p_pr, "spacing", after="240", before=before, line=LINE_SPACING,
        lineRule="auto")
    if align:
        sub(p_pr, "jc", val=align)
    elif justify:
        sub(p_pr, "jc", val="both")
    r_pr = sub(p_pr, "rPr")
    sub(r_pr, "sz", val=sz)
    sub(r_pr, "szCs", val=sz)
    if italic:
        runs = [(text, bold, True, mono, highlight)
                for text, bold, _, mono, highlight in runs]
    add_runs(paragraph, runs, sz=sz, color=color)
    return paragraph


def make_caption(runs):
    """The figure's alt text, under the picture: centred, small, grey."""
    return make_paragraph(runs, sz=CAPTION_SZ, justify=False, align="center",
                          color=CAPTION_COLOR, italic=True)


def make_heading(level, runs, page_break=False):
    paragraph = el("p")
    p_pr = sub(paragraph, "pPr")
    sub(p_pr, "pStyle", val=f"Heading{min(level, 6)}")
    if page_break:
        # The master sets pageBreakBefore on the FIRST Heading1 only, so the
        # body starts on a clean page after the contents list and every later
        # section flows normally. Setting it on all of them costs a blank page
        # per section.
        sub(p_pr, "pageBreakBefore", val="1")
    # Heading3 carries ind left=2160 hanging=360 in this template. Clearing it
    # needs both left=0 AND hanging=0; the firstLine=0 idiom the template uses
    # for H1/H2 is not enough, because those have no inherited indent to clear.
    sub(p_pr, "ind", left="0", right="0", hanging="0", firstLine="0")
    r_pr = sub(p_pr, "rPr")
    sub(r_pr, "color", val=HEAD_COLOR)
    add_runs(paragraph, runs, sz=None, color=HEAD_COLOR)
    return paragraph


def make_list_item(runs, num_id, level):
    paragraph = el("p")
    p_pr = sub(paragraph, "pPr")
    num_pr = sub(p_pr, "numPr")
    sub(num_pr, "ilvl", val=str(level))
    sub(num_pr, "numId", val=str(num_id))
    # The master sets list items to zero spacing, which runs wrapped bullets
    # together into a block. 3pt is enough to separate items without opening
    # the list up into something that reads as separate paragraphs.
    sub(p_pr, "spacing", after=LIST_ITEM_SPACING, before="0",
        line=LINE_SPACING, lineRule="auto")
    sub(p_pr, "ind", left=str(IND_LEFT[min(level, len(IND_LEFT) - 1)]),
        hanging=str(HANGING))
    sub(p_pr, "jc", val="both")
    r_pr = sub(p_pr, "rPr")
    sub(r_pr, "sz", val=BODY_SZ)
    sub(r_pr, "szCs", val=BODY_SZ)
    add_runs(paragraph, runs)
    return paragraph


def make_cell_paragraph(runs, header):
    paragraph = el("p")
    p_pr = sub(paragraph, "pPr")
    sub(p_pr, "keepNext", val="1")
    sub(p_pr, "keepLines", val="1")
    sub(p_pr, "spacing", after="0", before="0", lineRule="auto")
    if header:
        sub(p_pr, "jc", val="center")
    r_pr = sub(p_pr, "rPr")
    set_fonts(r_pr)
    if header:
        sub(r_pr, "b", val="1")
        sub(r_pr, "bCs", val="1")
    sub(r_pr, "sz", val=BODY_SZ)
    sub(r_pr, "szCs", val=BODY_SZ)

    for text, bold, italic, mono, highlight in runs:
        text_run(paragraph, text, bold=bold or header, italic=italic,
                 mono=mono, highlight=highlight, sz=BODY_SZ, font=BODY_FONT)
    if not paragraph.findall(qn("w:r")):
        sub(paragraph, "r")
    return paragraph


# Rough advance width of one character of 12pt Calibri, in twips, plus the
# cell padding. Only used to stop a column being narrower than its longest
# unbreakable word, so an approximation is enough.
# 3pt between list items. The master uses zero; see make_list_item.
LIST_ITEM_SPACING = "60"

TWIPS_PER_CHAR = 118
# Word's default cell margins are 108 twips a side; the rest is slack so a
# word does not sit hard against the border.
CELL_PADDING = 260
MAX_FLOOR = 2400
# A very long word is allowed to break rather than reserve the whole page;
# past about ten characters the reservation costs the prose columns too much.
MAX_FLOOR_CHARS = 10


def _column_text(rows, index):
    for row in rows:
        if index < len(row):
            yield "".join(run[0] for run in row[index])


def column_widths(rows, column_count):
    """Share the page width out in proportion to how much text each column holds.

    Dividing the width equally looks fine until one column carries a sentence
    and the others carry a single word, at which point the sentence column
    becomes a tall thin ribbon. Weighting by the longest cell fixes that.

    The floor matters as much as the weight: a column narrower than its longest
    word makes Word break the word itself, so a date column renders as "11/0
    8/20 26". Each column therefore reserves at least the width of its longest
    unbreakable token before the proportional share is applied.
    """
    min_weight, max_weight = 8, 60

    weights, floors = [], []
    for index in range(column_count):
        longest_cell = 0
        longest_word = 0
        for text in _column_text(rows, index):
            longest_cell = max(longest_cell, len(text))
            for word in text.split():
                longest_word = max(longest_word, len(word))
        weights.append(min(max(longest_cell, min_weight), max_weight))
        reserved = min(longest_word, MAX_FLOOR_CHARS)
        floors.append(min(CELL_PADDING + TWIPS_PER_CHAR * reserved, MAX_FLOOR))

    # A table of long words can want more than the page; scale the floors down
    # together rather than letting one column win.
    if sum(floors) > USABLE_TWIPS:
        scale = USABLE_TWIPS / sum(floors)
        floors = [int(f * scale) for f in floors]

    total = sum(weights)
    widths = [int(USABLE_TWIPS * w / total) for w in weights]

    # Raise anything below its floor, and pay for it from the columns that have
    # room to spare, in proportion to how much spare they have.
    deficit = sum(max(0, f - w) for w, f in zip(widths, floors))
    if deficit:
        surplus = [max(0, w - f) for w, f in zip(widths, floors)]
        total_surplus = sum(surplus)
        widths = [
            max(w, f) - (int(deficit * s / total_surplus) if total_surplus else 0)
            for w, f, s in zip(widths, floors, surplus)
        ]

    drift = sum(widths) - USABLE_TWIPS
    if drift:
        widest = widths.index(max(widths))
        widths[widest] -= drift
    return widths


def make_table(rows, style="Table4"):
    """Build a table using the template's direct-formatting recipe.

    The template's Table1..Table5 styles are functionally empty: they carry no
    borders and no fill, and every visible line and shade in the master is
    direct formatting in document.xml. Relying on the style would produce an
    invisible table, so the recipe is reproduced per cell.
    """
    if not rows:
        raise BodyError("cannot build a table with no rows")
    column_count = max(len(row) for row in rows)
    widths = column_widths(rows, column_count)

    table = el("tbl")
    tbl_pr = sub(table, "tblPr")
    sub(tbl_pr, "tblStyle", val=style)
    sub(tbl_pr, "tblW", w=str(USABLE_TWIPS), type="dxa")
    sub(tbl_pr, "jc", val="left")
    sub(tbl_pr, "tblInd", w="0", type="dxa")
    borders = sub(tbl_pr, "tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        sub(borders, side, color=TBL_BORDER, space="0", sz="4", val="single")
    # The master sets no cell margin at all, so descenders sit on the bottom
    # border and a table reads as a wall. 3pt top and bottom is the single
    # cheapest change to how a table feels.
    cell_margins = sub(tbl_pr, "tblCellMar")
    for side, width in (("top", CELL_MARGIN_V), ("left", CELL_MARGIN_H),
                        ("bottom", CELL_MARGIN_V), ("right", CELL_MARGIN_H)):
        sub(cell_margins, side, w=width, type="dxa")
    sub(tbl_pr, "tblLayout", type="fixed")
    sub(tbl_pr, "tblLook", val="0400")

    grid = sub(table, "tblGrid")
    for width in widths:
        sub(grid, "gridCol", w=str(width))

    for row_index, row in enumerate(rows):
        header = row_index == 0
        tr = sub(table, "tr")
        tr_pr = sub(tr, "trPr")
        # The master allows a row to break across a page, which strands half a
        # cell of prose at the top of the next page under no header.
        sub(tr_pr, "cantSplit", val="1")
        sub(tr_pr, "trHeight", val="390" if header else "300", hRule="atLeast")
        sub(tr_pr, "tblHeader", val="1" if header else "0")
        # Shading every body row the same makes a wide table hard to track
        # across. Banding alternate rows does the tracking for the reader, and
        # leaves the header band as the only strong horizontal.
        fill = HDR_FILL if header else (ROW_FILL if row_index % 2 == 0 else NO_FILL)
        for column_index in range(column_count):
            tc = sub(tr, "tc")
            tc_pr = sub(tc, "tcPr")
            sub(tc_pr, "tcW", w=str(widths[column_index]), type="dxa")
            cell_borders = sub(tc_pr, "tcBorders")
            for side in ("top", "left", "bottom", "right"):
                sub(cell_borders, side, color=CELL_BORDER, space="0",
                    sz="4", val="single")
            sub(tc_pr, "shd", fill=fill, val="clear")
            cell_runs = row[column_index] if column_index < len(row) else []
            tc.append(make_cell_paragraph(cell_runs, header))
    return table


# ------------------------------------------------------------------- pandoc --
def parse_markdown(markdown_text):
    try:
        result = subprocess.run(
            [find_pandoc(), "-f", PANDOC_FORMAT, "-t", "json"],
            input=markdown_text, capture_output=True, text=True, check=True,
        )
    except PandocNotFound as exc:
        raise BodyError(str(exc)) from exc
    except FileNotFoundError as exc:
        raise BodyError("pandoc could not be run. Check that the file it resolved "
                        "to is executable.") from exc
    except subprocess.CalledProcessError as exc:
        raise BodyError(f"pandoc could not parse the Markdown: {exc.stderr}") from exc
    return json.loads(result.stdout)


def table_rows(content):
    """Extract cell runs from a pandoc 2.10+ Table node."""
    _attr, _caption, _colspecs, thead, tbodies, tfoot = content
    rows = []

    def collect(row_list):
        for row in row_list:
            cells = []
            for cell in row[1]:
                runs = []
                for block in cell[4]:
                    if block["t"] in ("Plain", "Para"):
                        runs.extend(inline_runs(block["c"]))
                cells.append(runs)
            rows.append(cells)

    collect(thead[1])
    for body in tbodies:
        collect(body[3])
    if tfoot[1]:
        collect(tfoot[1])
    return rows


def lone_image(nodes):
    """Return (target, alt inlines) when `nodes` is one image and only that.

    An image on its own line is a figure. One sitting inside a sentence is not,
    and is left to the normal inline path, which renders its alt text.
    """
    if not isinstance(nodes, list):
        return None
    images = [n for n in nodes
              if isinstance(n, dict) and n.get("t") == "Image"]
    if len(images) != 1:
        return None
    rest = [n for n in nodes
            if isinstance(n, dict) and n.get("t") not in ("Image", "Space", "SoftBreak")]
    if rest:
        return None
    _attr, alt, target = images[0]["c"]
    return target[0], alt


def walk(blocks, emit, numbering, numberer, level=0, num_id=None, state=None):
    if state is None:
        state = {"first_heading": True}
    for block in blocks:
        tag, content = block["t"], block.get("c")

        figure = lone_image(content) if tag in ("Para", "Plain") else None
        if figure:
            target, alt = figure
            place_image = state.get("place_image")
            if place_image is None:
                raise BodyError(
                    f"the document has an image ({target}) but the renderer was "
                    "given no base directory to resolve it against")
            emit(place_image(target))
            alt_runs = inline_runs(alt)
            if alt_runs:
                emit(make_caption(alt_runs))

        elif tag == "Figure":
            # Pandoc 3 wraps a standalone image in a Figure. Its caption is the
            # same alt text, so the inner blocks alone carry everything.
            walk(content[2], emit, numbering, numberer, level, num_id, state)

        elif tag == "Header":
            heading_level, _attr, inlines = content
            runs = inline_runs(inlines)
            plain = "".join(r[0] for r in runs)
            prefix = numberer.prefix(heading_level, plain)
            if prefix:
                runs = [(prefix, False, False, False, None)] + runs
            # The break goes on the first heading whatever its level, so the
            # body always starts on a clean page after the contents list. Keyed
            # on level 1 alone it silently does nothing for a document whose
            # top heading is "##".
            page_break = state["first_heading"]
            if page_break:
                state["first_heading"] = False
            emit(make_heading(numberer.style_level(heading_level), runs,
                              page_break=page_break))

        elif tag in ("Para", "Plain"):
            runs = inline_runs(content)
            if level == 0:
                # A table carries no space beneath it, so the next paragraph
                # would otherwise sit hard against its bottom border.
                before = "240" if state.get("after_table") else "0"
                emit(make_paragraph(runs, before=before))
            else:
                emit(make_list_item(runs, num_id, min(level - 1, MAX_LIST_LEVEL)))

        elif tag == "BulletList":
            new_id = numbering.new(ordered=False)
            for item in content:
                walk(item, emit, numbering, numberer, level + 1, new_id, state)

        elif tag == "OrderedList":
            start = content[0][0] if content[0] else 1
            new_id = numbering.new(ordered=True, start=start)
            for item in content[1]:
                walk(item, emit, numbering, numberer, level + 1, new_id, state)

        elif tag == "Table":
            emit(make_table(table_rows(content)))

        elif tag == "DefinitionList":
            for term, definitions in content:
                emit(make_paragraph(inline_runs(term)))
                for definition in definitions:
                    walk(definition, emit, numbering, numberer, level, num_id, state)

        elif tag == "BlockQuote":
            walk(content, emit, numbering, numberer, level, num_id, state)

        elif tag == "Div":
            walk(content[1], emit, numbering, numberer, level, num_id, state)

        elif tag == "CodeBlock":
            emit(make_paragraph([(content[1], False, False, True, None)],
                                justify=False))

        elif tag == "HorizontalRule":
            emit(make_paragraph([]))

        state["after_table"] = tag == "Table"


def image_placer(doc, base_dir):
    """Add a picture to `doc` and hand back its detached paragraph.

    python-docx owns the media part and the relationship, so the picture is
    added through it and the resulting paragraph is then unhooked from the end
    of the body for `emit` to place in document order.
    """
    def place_image(target):
        path = (Path(base_dir) / target).expanduser()
        if not path.is_file():
            raise BodyError(f"image not found: {target}, looked in {base_dir}")
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        try:
            run.add_picture(str(path))
        except Exception as exc:
            raise BodyError(f"the image could not be embedded, {target}: {exc}") from exc
        shape = doc.inline_shapes[-1]
        if shape.width > IMAGE_MAX_EMU:
            shape.height = Emu(round(shape.height * IMAGE_MAX_EMU / shape.width))
            shape.width = Emu(IMAGE_MAX_EMU)
        p_pr = paragraph._p.get_or_add_pPr()
        sub(p_pr, "jc", val="center")
        element = paragraph._p
        element.getparent().remove(element)
        return element

    return place_image


def render(doc, markdown_text, numbering, heading_numbering="auto", base_dir=None):
    """Splice the rendered body into `doc`, before its sectPr."""
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is None:
        raise BodyError("the shell has no sectPr, so the page setup is missing")

    counter = [0]

    def emit(element):
        sect_pr.addprevious(element)
        counter[0] += 1

    ast = parse_markdown(markdown_text)
    numberer = HeadingNumberer(
        enabled=heading_numbering == "auto",
        top_level=shallowest_heading_level(ast["blocks"]),
    )
    state = {"first_heading": True}
    if base_dir is not None:
        state["place_image"] = image_placer(doc, base_dir)
    walk(ast["blocks"], emit, numbering, numberer, state=state)
    return counter[0]
