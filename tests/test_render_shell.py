"""The shell a built document must always have: styles, numbering, tables, cover.

These checks came from the standalone skill's selftest.py, which printed 16
results from one script run. Only the ones that still hold offline are here.
The rest needed LibreOffice to lay the document out, and that program is gone.

Every check runs against one build of the bundled example. No credential, no
network, no external program.

Where a check guards a house-style constant it states the value as a literal,
never by importing the constant. See the design rule in CLAUDE.md: a test that
reads the constant it tests passes whatever that constant is set to.
"""

import re
import zipfile

import pytest
from docx import Document

from gdoc.render import build
from gdoc.render.profiles import DEFAULT_TEMPLATE, TEMPLATES_DIR, resolve

EXAMPLE = TEMPLATES_DIR / "altery-group-policy-v1.0" / "example.md"

# The usable text width in twips: the A4 page's 11906 less the two 1021 margins.
# It is how a table we generated is told apart from one the master shipped, whose
# tblW is absent. Used to select tables, not to assert a house style.
USABLE_WIDTH = 9864

# 2pt, the least top and bottom cell padding that keeps descenders off the
# border. A floor rather than the CELL_MARGIN_V value, so setting that constant
# to the master's zero fails this file instead of redefining what passing means.
MIN_CELL_MARGIN_V = 40

# The house grey on every generated cell border. Lowercase here because that is
# what body.py writes; the comparison below is case-insensitive anyway, because
# an application that rewrites the file may normalise the case either way.
HOUSE_GREY = "c9c9c9"


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("shell") / "example.docx"
    build(EXAMPLE, out)
    return out


def part(path, name):
    return zipfile.ZipFile(path).read(f"word/{name}").decode("utf8")


def paragraph_text(paragraph_xml):
    return "".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", paragraph_xml))


def headings(path):
    """(style, text) for every paragraph carrying a Heading style, in order."""
    found = []
    for paragraph in re.findall(r"<w:p\b.*?</w:p>", part(path, "document.xml"), re.S):
        style = re.search(r'<w:pStyle w:val="(Heading\d)"/>', paragraph)
        if style:
            found.append((style.group(1), paragraph_text(paragraph)))
    return found


def heading_style(path, level):
    """The whole w:style block for HeadingN out of styles.xml."""
    match = re.search(rf'<w:style [^>]*w:styleId="Heading{level}".*?</w:style>',
                      part(path, "styles.xml"), re.S)
    assert match, f"Heading{level} is not defined"
    return match.group(0)


def our_nums(path):
    """{numId: its w:num definition} for the list instances this build added.

    Told apart from the master's by what the master defines, not by the overrides
    the checks below assert, which would make the selection circular. The master
    ships numId 1 to 7 and no lvlOverride at all.
    """
    master = set(re.findall(r'<w:num w:numId="(\d+)"',
                            part(resolve(DEFAULT_TEMPLATE), "numbering.xml")))
    found = {}
    for match in re.finditer(r'<w:num w:numId="(\d+)".*?</w:num>',
                             part(path, "numbering.xml"), re.S):
        if match.group(1) not in master:
            found[match.group(1)] = match.group(0)
    assert found, "no list instance was added, so there is nothing to check"
    return found


def front_matter_block():
    """The bundled example's front matter, closing --- included.

    Reused rather than hand-written so a document built from it validates the
    same way the real one does: every cover field, the revision rows and the
    classification are the example's own.
    """
    text = EXAMPLE.read_text(encoding="utf-8")
    end = text.index("\n---\n", text.index("---") + 3)
    return text[: end + len("\n---\n")]


def build_with_body(tmp_path, name, body_markdown):
    source = tmp_path / f"{name}.md"
    source.write_text(front_matter_block() + "\n" + body_markdown, encoding="utf-8")
    out = tmp_path / f"{name}.docx"
    build(source, out)
    return out


def example_title():
    return re.search(r"^title: (.+)$", front_matter_block(), re.M).group(1).strip()


def test_the_shallowest_heading_becomes_heading1(tmp_path):
    """Dropping the title's H1 changes which level is shallowest.

    Real notes start `# Title` then `## Section`. build drops the H1 that repeats
    the title, so the shallowest heading left is `##`, and that must come out as
    Heading1. Otherwise the whole document is demoted one level: the top sections
    render at 14pt in Heading2 and nothing carries Heading1 at all.

    Asserting which heading text carries Heading1, not merely that Heading1 and
    Heading2 both appear somewhere. That weaker form passes before and after the
    level shift exists, so it guards nothing. Matched on the text ending, because
    a heading carries its numbering prefix ("1-Top Level").
    """
    title = example_title()

    kept = headings(build_with_body(
        tmp_path, "kept",
        "# Top Level\n\ntext\n\n## Second\n\ntext\n",
    ))
    first = next(text for style, text in kept if style == "Heading1")
    assert first.endswith("Top Level"), f"Heading1 went to {first!r}"
    assert any(style == "Heading2" and text.endswith("Second") for style, text in kept)

    dropped = headings(build_with_body(
        tmp_path, "dropped",
        f"# {title}\n\ntext\n\n## Top Level\n\ntext\n\n### Second\n\ntext\n",
    ))
    assert not any(text.endswith(title) for _style, text in dropped), \
        "the title heading printed twice"
    first = next(text for style, text in dropped if style == "Heading1")
    assert first.endswith("Top Level"), f"Heading1 went to {first!r}"
    assert any(style == "Heading2" and text.endswith("Second") for style, text in dropped)


def test_every_referenced_paragraph_style_resolves(built):
    """A dangling w:pStyle is the most destructive defect available here.

    LibreOffice discards the whole w:pPr around it, so every bullet and every
    number in the document silently disappears.
    """
    defined = set(re.findall(r'w:styleId="([^"]+)"', part(built, "styles.xml")))
    referenced = set(re.findall(r'<w:pStyle w:val="([^"]+)"/>',
                                part(built, "document.xml")))
    assert referenced, "no paragraph style is referenced at all"
    assert not referenced - defined, f"dangling styles: {sorted(referenced - defined)}"


def test_every_numbering_id_used_is_defined(built):
    """An id written into the body but never defined leaves the list unnumbered."""
    used = set(re.findall(r'<w:numId w:val="(\d+)"/>', part(built, "document.xml")))
    defined = set(re.findall(r'<w:num w:numId="(\d+)"', part(built, "numbering.xml")))
    assert used, "the example has lists, so some numId must be used"
    assert not used - defined, f"undefined numIds: {sorted(used - defined, key=int)}"

    for num_id, definition in our_nums(built).items():
        overrides = re.findall(r"<w:lvlOverride\b.*?</w:lvlOverride>", definition, re.S)
        assert len(overrides) == 9, \
            f"numId {num_id} has {len(overrides)} lvlOverride, wanted 9"
        for level, override in enumerate(overrides):
            assert f'w:ilvl="{level}"' in override, \
                f"numId {num_id} override {level} is out of order: {override}"
            assert "<w:startOverride" in override, \
                f"numId {num_id} level {level} would carry on from the list above"


def test_a_second_ordered_list_restarts_at_one(tmp_path):
    """Two ordered lists in one document must not run 1, 2, 3, 4.

    A bare new w:num pointing at the same abstractNum does not restart numbering.
    The second list carries on from the first, which this project has already seen
    on the page as 1, 2, 2. What forces the restart is a w:num of its own with an
    explicit startOverride, so that is what is asserted: the two lists do not share
    an id, and each id restarts level zero at 1.

    Structural, because the rendered digit does not exist in the file. Nothing here
    lays the document out, so 1, 2, 1, 2 can only be read off the numbering
    definitions the way Word and Google read it.
    """
    out = build_with_body(
        tmp_path, "two_lists",
        "# Lists\n\n1. first\n2. second\n\ntext between\n\n1. third\n2. fourth\n",
    )
    order = re.findall(r'<w:numId w:val="(\d+)"/>', part(out, "document.xml"))
    ours = our_nums(out)
    used_in_order = [n for i, n in enumerate(order) if n not in order[:i] and n in ours]
    assert len(used_in_order) == 2, f"expected two lists, got numIds {used_in_order}"

    for num_id in used_in_order:
        level_zero = re.search(r'<w:lvlOverride w:ilvl="0">.*?</w:lvlOverride>',
                               ours[num_id], re.S)
        assert level_zero, f"numId {num_id} does not override level zero"
        assert '<w:startOverride w:val="1"/>' in level_zero.group(0), \
            f"numId {num_id} does not restart at 1: {level_zero.group(0)}"


def test_heading_styles_carry_house_spacing_indent_and_keeps(built):
    """1.15 line, room beneath, flush left, and never stranded or split.

    Every value here is a literal, and must stay one. Do not tidy 276, 120 and
    240 into HEADING_LINE_SPACING, HEADING_SPACE_AFTER and HEADING_SPACE_BEFORE:
    an assertion that reads the constant it guards follows that constant wherever
    anyone moves it, back to the master's single spacing and zero space beneath
    included, and still passes. Nor is a loose `after > 0` enough. It only catches
    the spacing being deleted, not 120 quietly becoming 20. The literal is what
    makes editing the constant fail this test, which is the whole point.

    keepNext and keepLines cover normalise_heading_keeps, which the standalone
    selftest never checked at all.
    """
    for level in range(1, 7):
        style = heading_style(built, level)
        spacing = re.search(r"<w:spacing[^>]*/>", style)
        assert spacing, f"Heading{level} declares no spacing"
        assert 'w:line="276"' in spacing.group(0), \
            f"Heading{level} is not 1.15 line: {spacing.group(0)}"
        assert 'w:after="120"' in spacing.group(0), \
            f"Heading{level} wants 6pt beneath it: {spacing.group(0)}"
        assert 'w:before="240"' in spacing.group(0), \
            f"Heading{level} wants 12pt above it: {spacing.group(0)}"

        indent = re.search(r"<w:ind[^>]*/>", style)
        assert indent, f"Heading{level} declares no indent"
        assert 'w:left="0"' in indent.group(0), f"Heading{level}: {indent.group(0)}"
        assert 'w:hanging="0"' in indent.group(0), f"Heading{level}: {indent.group(0)}"

        assert "<w:keepNext" in style, f"Heading{level} can be stranded on a page foot"
        assert "<w:keepLines" in style, f"Heading{level} can split across pages"


def test_heading_emphasis_falls_as_depth_grows(built):
    """A deeper heading must never look stronger than its parent.

    Heading3 to Heading6 all land on the same 12pt because the template's size
    scale runs out at level 3, so emphasis is the only thing separating them.
    The master has it backwards: Heading3 regular, Heading4 and Heading6 bold.

    Bold on 3 alone, italic on 5 and 6, stated as literals rather than read out
    of HEADING_EMPHASIS. A check echoing that dict would pass again the moment
    the master's inverted weights were put back.
    """
    for level, bold_wanted, italic_wanted in (
        (3, True, False), (4, False, False), (5, False, True), (6, False, True)
    ):
        style = heading_style(built, level)
        run_props = re.search(r"<w:rPr>.*?</w:rPr>", style, re.S)
        block = run_props.group(0) if run_props else ""
        # LibreOffice normalises val="1" to a bare <w:b/> and val="0" to
        # val="false", so presence of the tag alone does not mean the flag is on.
        bold = bool(re.search(r'<w:b(/>| w:val="(1|true|on)")', block))
        italic = bool(re.search(r'<w:i(/>| w:val="(1|true|on)")', block))
        assert bold is bold_wanted, f"Heading{level} bold={bold}, wanted {bold_wanted}"
        assert italic is italic_wanted, \
            f"Heading{level} italic={italic}, wanted {italic_wanted}"
        assert '<w:sz w:val="24"' in block, \
            f"Heading{level} is not 12pt: {block}"


def test_generated_tables_are_padded_grey_and_unsplit(built):
    """Only the tables we build are checked, found by their width.

    The three front-matter tables come from the master and keep its own black
    gridlines, so they carry no tblW at the usable width and drop out here.
    """
    document = part(built, "document.xml")
    generated = 0
    for index, match in enumerate(re.finditer(r"<w:tbl>.*?</w:tbl>", document, re.S)):
        table = match.group(0)
        width = re.search(r'<w:tblW w:w="(\d+)"', table)
        if not width or int(width.group(1)) != USABLE_WIDTH:
            continue
        generated += 1

        margins = re.search(r"<w:tblCellMar>.*?</w:tblCellMar>", table, re.S)
        vertical = re.findall(r'<w:(?:top|bottom) w:w="(\d+)"',
                              margins.group(0) if margins else "")
        assert vertical, f"table {index} sets no cell padding"
        assert all(int(v) >= MIN_CELL_MARGIN_V for v in vertical), \
            f"table {index} descenders sit on the border: {vertical}"

        colours = {c.lower() for c in re.findall(
            r'w:color="(\w{6})"',
            "".join(re.findall(r"<w:tcBorders>.*?</w:tcBorders>", table, re.S)))}
        # Compared lowercase. The standalone selftest compared against an
        # uppercased constant while the writer emits lowercase, so its own
        # check could never pass.
        assert colours == {HOUSE_GREY}, \
            f"table {index} border colours: {sorted(colours)}"

        rows = table.count("<w:tr>") + table.count("<w:tr ")
        unsplit = table.count("<w:cantSplit/>") + len(
            re.findall(r'<w:cantSplit w:val="(1|true|on)"', table))
        assert unsplit >= rows, f"table {index} has {rows - unsplit} splittable rows"
    assert generated, "the example has tables, so some must have been generated"


def test_front_matter_is_filled_and_the_cover_is_clean(built):
    """The cover and control pages must carry the note's own data, and nothing else.

    Yellow highlight means "a human must fill this in" and red means "this is
    guidance". Either one left on an issued document is a defect the reader acts
    on. Both are only defects on text: a real build leaves empty runs that still
    carry the master's red and its highlight, and an empty run shows nothing. The
    standalone selftest guarded its highlight branch this way but not its colour
    branch, so its own check failed on paragraph 7 of every build.
    """
    doc = Document(built)

    owner = doc.tables[0].rows[0].cells[1].text.strip()
    assert owner, "the version control table was left empty"

    revisions = front_matter_block().count("  - version:")
    assert len(doc.tables[1].rows) - 1 == revisions, \
        f"{len(doc.tables[1].rows) - 1} revision rows, front matter names {revisions}"

    shaded = [row.cells[0].text.strip() for row in doc.tables[2].rows[1:]
              if re.search(r'<w:shd [^/]*w:fill="(?!auto)[0-9a-fA-F]{6}"',
                           row.cells[1]._tc.xml)]
    assert shaded == ["Internal (I)"], f"shaded classification rows: {shaded}"

    leftovers = []
    for paragraph in doc.paragraphs[:20]:
        for run in paragraph.runs:
            xml = run._element.xml
            if not run.text.strip():
                continue
            if "<w:highlight" in xml:
                leftovers.append(f"highlight on {run.text.strip()[:24]!r}")
            if re.search(r'<w:color w:val="(ff0000|f00|cc0000)"', xml, re.I):
                leftovers.append(f"red text {run.text.strip()[:24]!r}")
    assert not leftovers, f"placeholder marks left on the cover: {leftovers}"
